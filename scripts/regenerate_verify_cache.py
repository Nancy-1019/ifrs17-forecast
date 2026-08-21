import os, sys
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'ifrs17_core.settings')
import django
django.setup()

from openpyxl import load_workbook
from data_input.views import (
    _load_calc_cache, _build_system_verify_rows, _generate_comparison,
    _CALC_CACHE, _VERIFY_CACHE, _save_verify_cache,
    VERIFY_OUTPUT_SHEETS, FINANCIAL_REPORT_SHEETS,
    _parse_verify_sheet, _parse_financial_report_sheet, _serialize_value
)

EXCEL_PATH = r'D:\IFRS17预测模型\验证文件\IFRS17财务预测_改造版_开发.xlsx'

_load_calc_cache()
calc = _CALC_CACHE.get('result')
if not calc:
    print('No calc result found. Please run calculation first.')
    sys.exit(1)

system_rows_map = _build_system_verify_rows(calc)
print('system sheets:', list(system_rows_map.keys()))
for sn, rows in system_rows_map.items():
    print(sn, 'rows:', len(rows), 'keys:', list(rows[0].keys())[:6] if rows else [])

# Build verify_data from attachment
wb = load_workbook(EXCEL_PATH, read_only=True, data_only=True)
verify_data = {}
for sheet_name in VERIFY_OUTPUT_SHEETS:
    if sheet_name not in wb.sheetnames:
        verify_data[sheet_name] = {'available': False, 'headers': [], 'rows': [], 'totalRows': 0}
        continue
    ws = wb[sheet_name]
    if sheet_name in FINANCIAL_REPORT_SHEETS:
        headers, rows = _parse_financial_report_sheet(ws)
    else:
        headers, rows = _parse_verify_sheet(ws)
    rows_serialized = [[_serialize_value(v) for v in row] for row in rows]
    verify_data[sheet_name] = {
        'available': True,
        'headers': [_serialize_value(h) for h in headers],
        'rows': rows_serialized,
        'totalRows': len(rows_serialized),
    }
wb.close()

comparison = _generate_comparison(verify_data, system_rows_map=system_rows_map)
_VERIFY_CACHE['sheet_data'] = verify_data
_VERIFY_CACHE['comparison'] = comparison
_VERIFY_CACHE['file_name'] = os.path.basename(EXCEL_PATH)
from django.utils import timezone
_VERIFY_CACHE['timestamp'] = timezone.now().isoformat()
_save_verify_cache()

print('\n=== comparison result ===')
for sn, cmp in comparison.items():
    print(sn, 'diffFieldCount:', cmp.get('diffFieldCount'), 'totalDiffAmount:', cmp.get('totalDiffAmount'))
