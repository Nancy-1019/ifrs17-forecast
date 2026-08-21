"""Add the missing '提前确认保单_现有业务' sheet to the dock template.

Spec (SYSTEM_DOCK_SPECS) requires 10 dock sheets including
'提前确认保单_现有业务' (更新日期/评估时点/数据类型/精算险类 + 60月),
but the template only had 9. This makes template == spec sheet set.
"""
import os
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
TEMPLATE = os.path.join(ROOT, '验证文件', 'backup', '系统对接输入表_模板.xlsx')

wb = openpyxl.load_workbook(TEMPLATE)
print("模板现有 sheets:", wb.sheetnames)
if '提前确认保单_现有业务' not in wb.sheetnames:
    ws = wb.create_sheet(title='提前确认保单_现有业务')
    headers = ['更新日期', '评估时点', '数据类型', '精算险类'] + [str(i) for i in range(1, 61)]
    ws.append(headers)
    print(f"已新增 '提前确认保单_现有业务' ({len(headers)}列)")
else:
    print("已存在，跳过")
wb.save(TEMPLATE)
wb.close()
print("已保存:", TEMPLATE)
