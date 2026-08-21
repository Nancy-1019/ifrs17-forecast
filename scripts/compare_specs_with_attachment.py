# -*- coding: utf-8 -*-
"""对比当前系统输入表字段规范(excel_specs.py) 与 附件最新输入表字段的差异。"""
import sys, os
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'ifrs17-system')))
import openpyxl
from data_input.excel_specs import SYSTEM_DOCK_SPECS, EXCEL_UPLOAD_SPECS

PROJ = r"D:\IFRS17预测模型"
ATT = os.path.join(PROJ, "验证文件", "IFRS17财务预测_改造版_开发.xlsx")

ATT_HEADER_ROW = {"初始确认利率曲线": 3}
NAME_MAP = {"说明": "输入表说明"}


def read_headers(ws, row=1):
    out = []
    for c in range(1, ws.max_column + 1):
        out.append(ws.cell(row=row, column=c).value)
    while out and out[-1] is None:
        out.pop()
    return out


def trim_header(hdr):
    k = len(hdr)
    while k > 0 and not isinstance(hdr[k - 1], str):
        k -= 1
    return hdr[:k]


def main():
    wb = openpyxl.load_workbook(ATT, read_only=True, data_only=True)
    att_sheets = wb.sheetnames
    att_headers = {}
    for sn in att_sheets:
        hrow = ATT_HEADER_ROW.get(sn, 1)
        h = read_headers(wb[sn], hrow)
        if sn == "初始确认利率曲线":
            h = trim_header(h)
        att_headers[sn] = h
    wb.close()

    print("=" * 110)
    print("附件全部工作表 (共 %d 个):" % len(att_sheets))
    print(" | ".join(att_sheets))
    print("=" * 110)

    for spec_name, specs in [("系统对接输入表", SYSTEM_DOCK_SPECS), ("手工输入表", EXCEL_UPLOAD_SPECS)]:
        print("\n########## %s (共 %d 张) ##########" % (spec_name, len(specs)))
        for sn, sp in specs.items():
            a_sn = NAME_MAP.get(sn, sn)
            if a_sn not in att_headers:
                print(f"\n[缺失] 规范表 '{sn}' 在附件中无对应表(期望附件表名 '{a_sn}')")
                continue
            spec_set = [str(x) for x in sp.headers]
            att_set = [str(x) for x in att_headers[a_sn] if x is not None]
            spec_idx = {h: i for i, h in enumerate(spec_set)}
            att_idx = {h: i for i, h in enumerate(att_set)}
            missing = [h for h in spec_set if h not in att_idx]      # 规范有、附件无
            extra = [h for h in att_set if h not in spec_idx]        # 附件有、规范无
            # 顺序是否一致
            order_ok = (spec_set == att_set)
            flag = "OK" if (not missing and not extra and order_ok) else "DIFF"
            print(f"\n[{flag}] '{sn}' (附件表='{a_sn}') 规范列={len(spec_set)} 附件列={len(att_set)}"
                  + (" 顺序一致" if order_ok else " 顺序不一致"))
            if missing:
                print("   规范有但附件缺: " + ", ".join(missing))
            if extra:
                print("   附件有但规范缺: " + ", ".join(extra))
            if not order_ok and not missing and not extra:
                # 同集合不同序：找错位
                diffpos = [i for i in range(len(spec_set)) if spec_set[i] != att_set[i]]
                print("   前若干错位(规范->附件): " + "; ".join(
                    f"{spec_set[i]}->{att_set[i]}" for i in diffpos[:8]))


if __name__ == "__main__":
    main()
