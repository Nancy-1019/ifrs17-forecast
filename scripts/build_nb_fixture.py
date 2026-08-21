# -*- coding: utf-8 -*-
"""Build new-business input fixture from the authoritative verification workbook.

The oracle (scripts/vba_faithful_new.py) reads all new-business inputs directly from
IFRS17财务预测_改造版_0729_VBA.xlsx with specific key columns per sheet. The uploaded
手工输入表/系统对接输入表 copies of these pattern tables have inconsistent / empty key
columns (e.g. 未到期赚取模式_新业务 has an empty 预测组 column for every row), which makes
the engine's pattern lookups fail (L=当期确认比例=0 -> all 输出_* wrong).

To guarantee the engine receives IDENTICAL inputs to the proven oracle, this script
replicates the oracle's exact table_keyed / table_keyed_monthcols calls against the
validation workbook and serialises the resulting dicts to
data_input/fixtures/new_business_inputs.json.

Oracle keying (1-indexed key_col):
  keynew    输入整理-合同组关键假设_新业务  key_col=4 (预测组ID)
  ep_new    输入整理-预期赔付率            key_col=5 (预测组)
  maint_new 输入整理-维持费用率            key_col=5 (合同组ID)
  und_incur 输入整理-未到期赔付模式        key_col=4 (险类)
  cum       输入整理-实际赔付比例_累计      key_col=4 (险类)
  prem_mode 输入整理-保费现金流模式        key_col=4 (合同组ID)
  iacf_mode 输入整理-IACF现金流模式        key_col=4 (合同组ID)
  earn      输入整理-未到期赚取模式        key_col=4 (合同组ID)
  riskadj   风险调整比例                   key_col=3 (险类)
  reins     再保人不履约风险               key_col=3 (险类)
  invest    投资成分比例                   key_col=3 (预测组)
  recov_new 预期摊回比例_新业务            key_col=3 (险类)
  io_indir  未决间接理赔费用率             key_col=3 (险类)
"""
import os, json, datetime

ROOT = "D:/IFRS17预测模型"
VBX = os.path.join(ROOT, "验证文件", "IFRS17财务预测_改造版_0729_VBA.xlsx")
OUT = os.path.join(ROOT, "ifrs17-system", "data_input", "fixtures", "new_business_inputs.json")

from openpyxl import load_workbook
wb = load_workbook(VBX, data_only=True)


def norm_key(v):
    if isinstance(v, datetime.datetime):
        return str(v.toordinal() - datetime.date(1899, 12, 30).toordinal())
    return str(v)


def _norm_header(h):
    if h is None:
        return None
    return "".join(str(h).split())


def table_keyed(sheet, key_col, header_row=1):
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [_norm_header(h) for h in rows[header_row - 1]]
    ki = key_col - 1
    out = {}
    for r in rows[header_row:]:
        if ki >= len(r) or r[ki] is None:
            continue
        key = norm_key(r[ki])
        d = {}
        for j, v in enumerate(r):
            h = hdr[j]
            if h is not None and j != ki:
                d[h] = (None if v is None else (v if isinstance(v, (int, float)) else str(v)))
        out[key] = d
    return out


def table_keyed_monthcols(sheet, key_col, header_row=1):
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [_norm_header(h) for h in rows[header_row - 1]]
    ki = key_col - 1
    out = {}
    for r in rows[header_row:]:
        if ki >= len(r) or r[ki] is None:
            continue
        key = norm_key(r[ki])
        d = {}
        for j, v in enumerate(r):
            if j == ki:
                continue
            h = hdr[j]
            if h is None:
                continue
            try:
                m = int(float(h))
            except Exception:
                continue
            d[m] = (0.0 if v is None else float(v))
        out[key] = d
    return out


def to_float(v):
    try:
        return float(v) if v is not None else 0.0
    except Exception:
        return 0.0


# ---- keynew: keep only the fields the engine / oracle need ----
keynew_raw = table_keyed("输入整理-合同组关键假设_新业务", 4)
KEEP = ["保费收入", "跟单获取费用（净额结算手续费）", "非跟单获取费用",
        "当期提前初始确认签单保费", "评估期开始业务预期写入时间（月）",
        "业务类型", "精算险类", "预测组"]
keynew = {}
for gid, row in keynew_raw.items():
    rec = {}
    for f in KEEP:
        rec[f] = row.get(f)
    # 将原 Oracle/VBA 单字母字段 U/V/CC/E 映射为规范中英文字段名
    rec["premium_income"] = to_float(row.get("保费收入"))
    rec["acquisition_cost"] = to_float(row.get("跟单获取费用（净额结算手续费）")) + to_float(row.get("非跟单获取费用"))
    rec["advance_written_premium"] = to_float(row.get("当期提前初始确认签单保费"))
    rec["expected_writing_period_months"] = row.get("评估期开始业务预期写入时间（月）")
    keynew[gid] = rec

data = {
    "keynew": keynew,
    "earn": table_keyed_monthcols("输入整理-未到期赚取模式", 4),
    "prem_mode": table_keyed_monthcols("输入整理-保费现金流模式", 4),
    "iacf_mode": table_keyed_monthcols("输入整理-IACF现金流模式", 4),
    "ep": table_keyed_monthcols("输入整理-预期赔付率", 5),
    "maint": table_keyed_monthcols("输入整理-维持费用率", 5),
    "invest": table_keyed_monthcols("投资成分比例", 3),
    "und_incur": table_keyed_monthcols("输入整理-未到期赔付模式", 4),
    "cum": table_keyed_monthcols("输入整理-实际赔付比例_累计", 4),
    "riskadj": table_keyed("风险调整比例", 3),
    "reins": table_keyed_monthcols("再保人不履约风险", 3),
    "recov": table_keyed_monthcols("预期摊回比例_新业务", 3),
    "io_indir": table_keyed_monthcols("未决间接理赔费用率", 3),
}

# ---- discount curve + monthly interest rate (oracle source: 输入整理-即期利率曲线加工) ----
wsr = wb["输入整理-即期利率曲线加工"]
rowsr = list(wsr.iter_rows(values_only=True))
hdr_r = [_norm_header(h) for h in rowsr[0]]
ie = hdr_r.index("进展月份") if "进展月份" in hdr_r else None
ih = hdr_r.index("月度折现因子（期末折现）_预测时点") if "月度折现因子（期末折现）_预测时点" in hdr_r else None
ii = hdr_r.index("月度折现因子（期初折现）_预测时点") if "月度折现因子（期初折现）_预测时点" in hdr_r else None
ij = hdr_r.index("月度远期利率_预测时点") if "月度远期利率_预测时点" in hdr_r else None
disc_end, disc_begin, intrate = {}, {}, {}
for r in rowsr[1:]:
    if ie is None or r[ie] is None:
        continue
    try:
        k = int(float(r[ie]))
    except Exception:
        continue
    disc_end[k] = float(r[ih]) if ih is not None and r[ih] is not None else 0.0
    disc_begin[k] = float(r[ii]) if ii is not None and r[ii] is not None else 0.0
    intrate[k] = float(r[ij]) if ij is not None and r[ij] is not None else 0.0
# extend beyond max key by compounding last monthly factor (mirror oracle)
maxk = max(disc_end.keys()) if disc_end else 0
if maxk > 1:
    r_end = disc_end[maxk] / disc_end[maxk - 1] if disc_end[maxk - 1] != 0 else 1.0
    r_begin = disc_begin[maxk] / disc_begin[maxk - 1] if disc_begin[maxk - 1] != 0 else 1.0
    r_int = intrate[maxk] if intrate.get(maxk) else 0.0
    for k in range(maxk + 1, 125):
        disc_end[k] = disc_end[k - 1] * r_end
        disc_begin[k] = disc_begin[k - 1] * r_begin
        intrate[k] = r_int
data["disc_end"] = disc_end
data["disc_begin"] = disc_begin
data["intrate"] = intrate

os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", encoding="utf-8") as f:
    json.dump(data, f, ensure_ascii=False)

# ---- summary ----
print("Wrote", OUT)
print("keynew groups:", len(keynew))
for gid in list(keynew)[:3]:
    r = keynew[gid]
    print(f"  {gid}: premium_income={r['premium_income']} acquisition_cost={r['acquisition_cost']} "
          f"advance_written_premium={r['advance_written_premium']} expected_writing_period_months={r['expected_writing_period_months']} "
          f"biz={r['业务类型']} 险类={r['精算险类']} 预测组={r['预测组']}")
for k in ["earn", "prem_mode", "iacf_mode", "ep", "maint", "invest", "und_incur", "cum", "riskadj", "reins", "recov", "io_indir"]:
    d = data[k]
    sample_key = next(iter(d)) if d else None
    nm = len(d.get(sample_key, {})) if sample_key else 0
    print(f"  {k}: groups={len(d)} sample_key={sample_key} n_month_or_field={nm}")
print(f"  disc_end: {len(disc_end)} (maxk={max(disc_end)})  intrate: {len(intrate)}")
