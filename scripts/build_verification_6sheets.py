import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

src = 'D:/IFRS17预测模型/验证文件/IFRS17财务预测_改造版_0729_VBA.xlsx'
dst = 'D:/IFRS17预测模型/验证文件/验证核对工作表_6张.xlsx'

keep_sheets = ['PAA计算_新业务整理', 'PAA计算_现有业务整理', 'PAA计算_汇总', 'PAA计算_MTD', '输出财务报表_MTD', '输出财务报表_YTD']

print("Loading source workbook...")
wb = openpyxl.load_workbook(src, data_only=False)

missing = [s for s in keep_sheets if s not in wb.sheetnames]
if missing:
    print("MISSING SHEETS:", missing)
    raise SystemExit(1)

# Remove all sheets not in keep list
to_remove = [s for s in wb.sheetnames if s not in keep_sheets]
print(f"Removing {len(to_remove)} unwanted sheets...")
for s in to_remove:
    print(f"  Remove: {s}")
    wb.remove(wb[s])

# Reorder sheets to match keep_sheets order
# openpyxl indexes sheets; create ordered references
ordered = [wb[s] for s in keep_sheets]
wb._sheets = ordered

# Insert summary sheet at front
ws_summary = wb.create_sheet("核对说明", 0)
ws_summary.column_dimensions['A'].width = 28
ws_summary.column_dimensions['B'].width = 60
ws_summary.column_dimensions['C'].width = 18

header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
header_font = Font(color="FFFFFF", bold=True, size=12)
thin_border = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin')
)

headers = ['工作表名称', '用途说明', '来源']
for col, h in enumerate(headers, 1):
    cell = ws_summary.cell(row=1, column=col, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = thin_border

descriptions = {
    'PAA计算_新业务整理': '新业务 PAA 计量结果：包含 41 组合同组的保费、获取费用、未到期责任负债、已发生未决赔款负债、保险合同收入等输出列，用于与系统 newBusinessPredict 比对。',
    'PAA计算_现有业务整理': '现有业务 PAA 计量结果：包含 41 组合同组的期初/期末现值、待摊销金额、责任负债、现金流等输出列，用于与系统 existingBusinessPredict 比对。',
    'PAA计算_汇总': '新老业务汇总结果：按合同组/险种维度聚合 PAA 输出科目，用于与系统汇总表比对。',
    'PAA计算_MTD': '月度累计（MTD）视角的 PAA 结果，用于校验当月损益与资产负债变动。',
    '输出财务报表_MTD': 'MTD 损益表与资产负债表输出，用于校验系统财务报表 MTD 结果。',
    '输出财务报表_YTD': 'YTD 损益表与资产负债表输出，用于校验系统财务报表 YTD 结果。',
}

for row, sheet_name in enumerate(keep_sheets, 2):
    ws_summary.cell(row=row, column=1, value=sheet_name).border = thin_border
    ws_summary.cell(row=row, column=2, value=descriptions[sheet_name]).border = thin_border
    ws_summary.cell(row=row, column=3, value='0729_VBA.xlsx').border = thin_border
    for col in range(1, 4):
        ws_summary.cell(row=row, column=col).alignment = Alignment(vertical='center', wrap_text=True)

note_start = len(keep_sheets) + 3
notes = [
    ('生成时间', '2026-07-31'),
    ('生成依据', '验证文件/IFRS17财务预测_改造版_0729_VBA.xlsx'),
    ('使用方式', '与系统计量输出（系统计量输出_验证格式.xlsx 或 runtime cache）按共同字段逐行比对；数值差异应控制在容差 0.01 以内。'),
]
for i, (k, v) in enumerate(notes, note_start):
    ws_summary.cell(row=i, column=1, value=k).font = Font(bold=True)
    ws_summary.cell(row=i, column=2, value=v).alignment = Alignment(wrap_text=True)

ws_summary.freeze_panes = 'A2'

print(f"Saving to {dst}...")
wb.save(dst)
print(f"Saved verification workbook: {dst}")
print(f"Sheets: {wb.sheetnames}")
