# -*- coding: utf-8 -*-
import os
from openpyxl import load_workbook
ROOT="D:/IFRS17预测模型"
VBX=os.path.join(ROOT,"验证文件","IFRS17财务预测_改造版_0729_VBA.xlsx")
wb=load_workbook(VBX,data_only=True)
ws=wb["现有业务预期现金流_1"]
hdr=[c.value for c in ws[1]]
agi=hdr.index("事故期月度间隔"); ahi=hdr.index("发展期月度间隔")
print("AGidx",agi,"AHidx",ahi)
print("=== rows 2..125 (row, a=AG, d=AH) ===")
for ri in range(2,126):
    ag=ws.cell(row=ri,column=agi+1).value
    ah=ws.cell(row=ri,column=ahi+1).value
    print(f" row{ri}: a={ag} d={ah}")
wb.close()
