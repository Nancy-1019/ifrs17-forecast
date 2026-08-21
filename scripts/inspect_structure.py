# -*- coding: utf-8 -*-
import os
from openpyxl import load_workbook
ROOT="D:/IFRS17预测模型"
VBX=os.path.join(ROOT,"验证文件","IFRS17财务预测_改造版_0729_VBA.xlsx")
wb=load_workbook(VBX,data_only=True)
print("SHEETS:", [s for s in wb.sheetnames if "预期现金流" in s or "PAA计算" in s][:80])
print()
# 预测 table rows: D=合同组ID, J=预测间隔, K=数据来源表, C=合同组名称, E=直保/分出
ws=wb["PAA计算_现有业务预测"]
rows=list(ws.iter_rows(values_only=True))
hdr=rows[0]
def ci(n): return hdr.index(n)
print("预测 table: rows=",len(rows)-1)
seen={}
for r in rows[1:]:
    gid=r[ci("合同组ID")]; j=r[ci("预测间隔")]; k=r[ci("数据来源表")]; c=r[ci("合同组名称")]; e=r[ci("直保或分入/分出")]
    seen.setdefault(gid,[]).append((j,k,e))
for gid,vals in seen.items():
    print(f"  {gid}: n={len(vals)} e.g. {vals[0]}")
print()
# 整理 table
wz=wb["PAA计算_现有业务整理"]
rz=list(wz.iter_rows(values_only=True))
hz=rz[0]
print("整理 table rows=",len(rz)-1)
zgid=set()
for r in rz[1:]:
    zgid.add((r[hz.index("合同组ID")], r[hz.index("直保或分入/分出")]))
print("整理 unique (gid,biz):",len(zgid))
for x in list(zgid)[:10]: print("   ",x)
wb.close()
