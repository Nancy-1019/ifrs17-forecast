#!/usr/bin/env python
import os
from openpyxl import load_workbook

OUT = r'D:\IFRS17预测模型\验证文件\test_export\验证核对完整结果.xlsx'
wb = load_workbook(OUT, read_only=True, data_only=True)
print('sheets:', wb.sheetnames)
ws = wb['差异汇总']
print('\n--- 差异汇总 ---')
for row in ws.iter_rows(values_only=True):
    print(row)
wb.close()

# 检查 汇总 表头与样本（验证/系统/差异）
wb = load_workbook(OUT, read_only=True, data_only=True)
ws = wb['PAA计算_汇总']
hdr = [c.value for c in next(ws.iter_rows(min_row=1, max_row=1))]
print('\n汇总 headers (first 12):', hdr[:12], '... total', len(hdr))
print('汇总 sample rows (first 3):')
for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i == 0: continue
    if i > 3: break
    print([row[j] for j in range(min(12, len(row)))])
wb.close()

# 检查 comparison diffFieldCount（从 verify_cache.json）
import json
vc = json.load(open(r'D:\IFRS17预测模型\ifrs17-system\data_input\runtime_cache\verify_cache.json', encoding='utf-8'))
cmp = vc.get('comparison') or {}
print('\n--- comparison per sheet ---')
for sn, c in cmp.items():
    print(sn, 'status=', c.get('status'), 'diffFieldCount=', c.get('diffFieldCount'), 'totalDiff=', c.get('totalDiffAmount'))
