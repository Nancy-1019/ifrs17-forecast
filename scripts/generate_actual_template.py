"""
生成预实分析实际数据上传模板
路径: ifrs17-system/scripts/generate_actual_template.py
输出: ifrs17-system/验证文件/预实分析_实际数据上传模板.xlsx
"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, NamedStyle
from openpyxl.utils import get_column_letter

OUT_DIR = os.path.join(os.path.dirname(__file__), '..', '验证文件')
OUT_FILE = os.path.join(OUT_DIR, '预实分析_实际数据上传模板.xlsx')

# 精算险类代码 -> 名称（与前端 ACTUARIAL_CLASS_NAMES 对齐）
ACTUARIAL_CLASSES = [
    ('32', '交强险'),
    ('36', '附加险'),
    ('37', '商业三者险'),
    ('38', '车损险'),
    ('101', '企业财产险'),
    ('102', '货运险'),
    ('104', '责任险'),
    ('105', '保证保险'),
    ('106', '信用保险'),
    ('107', '特殊风险'),
    ('108', '家财险'),
    ('109', '建工险'),
    ('110', '其他险'),
    ('111', '船舶险'),
    ('301', '种植业'),
    ('302', '养殖业'),
    ('602', '意外险'),
    ('603CC', '商业性健康险'),
    ('603ZD', '大病保险'),
    ('603ZX', '其他政策性健康险'),
    ('999', '其他'),
]

# 公司合计实际数科目（与现有 api_upload_actual 解析逻辑对齐）
COMPANY_ITEMS = [
    '保险服务收入',
    '保险服务费用',
    '分出再保险净损益',
    '承保财务损益净额',
    '投资收益',
    '其他损益',
    '净利润',
]

# 精算险类实际数列（金额类 + 比率类）
CLASS_COLUMNS = [
    '评估时点',
    '精算险类代码',
    '精算险类名称',
    '保险服务收入',
    '保险服务费用',
    '分出再保险净损益',
    '承保财务损益净额',
    '投资收益',
    '其他损益',
    '营业利润',
    '净利润',
    '净资产',
    '综合成本率(%)',
    '综合赔付率(%)',
    '综合费用率(%)',
]

# 样式定义
header_fill = PatternFill(start_color='E6F4FF', end_color='E6F4FF', fill_type='solid')
header_font = Font(name='Microsoft YaHei', size=11, bold=True, color='1677FF')
subheader_fill = PatternFill(start_color='F5F5F5', end_color='F5F5F5', fill_type='solid')
subheader_font = Font(name='Microsoft YaHei', size=10, bold=True)
normal_font = Font(name='Microsoft YaHei', size=10)
thin_border = Border(
    left=Side(style='thin', color='D9D9D9'),
    right=Side(style='thin', color='D9D9D9'),
    top=Side(style='thin', color='D9D9D9'),
    bottom=Side(style='thin', color='D9D9D9')
)

os.makedirs(OUT_DIR, exist_ok=True)
wb = Workbook()

# ============================================================
# 工作表1：填报说明
# ============================================================
ws1 = wb.active
ws1.title = '填报说明'
ws1.sheet_properties.tabColor = '1677FF'

instructions = [
    ['预实分析实际数据上传模板', ''],
    ['', ''],
    ['一、模板用途', ''],
    ['本模板用于向 IFRS17 预测模型系统上传“实际”财务结果，与系统生成的“预测/预算”结果进行比对。', ''],
    ['支持两个维度：', ''],
    ['  1. 公司合计实际数：用于整体预实偏差分析、达成率看板。', ''],
    ['  2. 精算险类实际数：用于分险种（精算险类）预实对比分析。', ''],
    ['', ''],
    ['二、工作表说明', ''],
    ['工作表名称', '说明'],
    ['公司合计实际数', '每行一个科目，填写该科目的实际发生额。评估时点通常为月度/季度末，如 2026-06-30。'],
    ['精算险类实际数', '每行一个精算险类，填写该公司各险种科目的实际发生额。'],
    ['精算险类代码对照', '精算险类代码与中文名称对照表，供复制使用，无需修改。'],
    ['', ''],
    ['三、填报规则', ''],
    ['1. 评估时点', '格式建议 YYYY-MM-DD，例如 2026-06-30。公司合计与精算险类可分别填报不同时点。'],
    ['2. 金额类科目', '以人民币“元”为单位，可直接填写数值，不需要千分位分隔符，支持小数。'],
    ['3. 比率类科目', '综合成本率/综合赔付率/综合费用率请按百分比填写，如 98.21 表示 98.21%。'],
    ['4. 精算险类代码', '请填写“精算险类代码对照”工作表中的标准代码，确保与预测结果中的精算险类字段一致。'],
    ['5. 空值处理', '不涉及的科目可留空，系统将按 0 处理。'],
    ['', ''],
    ['四、科目口径', ''],
    ['保险服务收入', '直保/分入业务在 PAA 下的保险合同收入合计。'],
    ['保险服务费用', '已发生赔款理赔费用、获取费用摊销、维持费用等保险服务费用合计。'],
    ['分出再保险净损益', '分出保费的分摊 - 摊回保险服务费用 + 分出再保险承保财务损失等净额。'],
    ['承保财务损益净额', '承保财务损失/收益净额。'],
    ['投资收益', '利润表中的投资收益科目。'],
    ['其他损益', '营业外收支、利息收支、信用减值损失等其他损益净额。'],
    ['营业利润', '承保利润 + 投资收益 + 其他损益。'],
    ['净利润', '利润总额 - 所得税费用。'],
    ['净资产', '资产负债表中的所有者权益合计。'],
]

for r_idx, (col1, col2) in enumerate(instructions, 1):
    ws1.cell(row=r_idx, column=1, value=col1)
    ws1.cell(row=r_idx, column=2, value=col2)
    if r_idx == 1:
        ws1.cell(row=r_idx, column=1).font = Font(name='Microsoft YaHei', size=16, bold=True, color='1677FF')
    elif col1 in ['一、模板用途', '二、工作表说明', '三、填报规则', '四、科目口径']:
        ws1.cell(row=r_idx, column=1).font = Font(name='Microsoft YaHei', size=12, bold=True, color='262626')
    elif col1 == '工作表名称':
        ws1.cell(row=r_idx, column=1).font = subheader_font
        ws1.cell(row=r_idx, column=2).font = subheader_font
        ws1.cell(row=r_idx, column=1).fill = subheader_fill
        ws1.cell(row=r_idx, column=2).fill = subheader_fill

ws1.column_dimensions['A'].width = 28
ws1.column_dimensions['B'].width = 90

# ============================================================
# 工作表2：公司合计实际数
# ============================================================
ws2 = wb.create_sheet('公司合计实际数')
ws2.sheet_properties.tabColor = '52C41A'

ws2.cell(row=1, column=1, value='评估时点')
ws2.cell(row=1, column=2, value='科目')
ws2.cell(row=1, column=3, value='实际值')
ws2.cell(row=1, column=4, value='说明')

for c in range(1, 5):
    cell = ws2.cell(row=1, column=c)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = thin_border

example_date = '2026-06-30'
for i, item in enumerate(COMPANY_ITEMS, 2):
    ws2.cell(row=i, column=1, value=example_date)
    ws2.cell(row=i, column=2, value=item)
    ws2.cell(row=i, column=3, value=None)
    ws2.cell(row=i, column=4, value='金额，单位：元')
    for c in range(1, 5):
        ws2.cell(row=i, column=c).border = thin_border
        ws2.cell(row=i, column=c).font = normal_font

ws2.column_dimensions['A'].width = 16
ws2.column_dimensions['B'].width = 24
ws2.column_dimensions['C'].width = 20
ws2.column_dimensions['D'].width = 20

# ============================================================
# 工作表3：精算险类实际数
# ============================================================
ws3 = wb.create_sheet('精算险类实际数')
ws3.sheet_properties.tabColor = 'FAAD14'

for c_idx, col_name in enumerate(CLASS_COLUMNS, 1):
    cell = ws3.cell(row=1, column=c_idx, value=col_name)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border

example_date = '2026-06-30'
for r_idx, (code, name) in enumerate(ACTUARIAL_CLASSES, 2):
    ws3.cell(row=r_idx, column=1, value=example_date)
    ws3.cell(row=r_idx, column=2, value=code)
    ws3.cell(row=r_idx, column=3, value=name)
    for c in range(4, len(CLASS_COLUMNS) + 1):
        ws3.cell(row=r_idx, column=c, value=None)
    for c in range(1, len(CLASS_COLUMNS) + 1):
        ws3.cell(row=r_idx, column=c).border = thin_border
        ws3.cell(row=r_idx, column=c).font = normal_font

for c_idx, col_name in enumerate(CLASS_COLUMNS, 1):
    if col_name in ['精算险类名称']:
        ws3.column_dimensions[get_column_letter(c_idx)].width = 22
    elif col_name in ['评估时点', '精算险类代码']:
        ws3.column_dimensions[get_column_letter(c_idx)].width = 14
    elif '率' in col_name:
        ws3.column_dimensions[get_column_letter(c_idx)].width = 16
    else:
        ws3.column_dimensions[get_column_letter(c_idx)].width = 18

# 固定表头
ws3.freeze_panes = 'A2'

# ============================================================
# 工作表4：精算险类代码对照
# ============================================================
ws4 = wb.create_sheet('精算险类代码对照')
ws4.sheet_properties.tabColor = '8C8C8C'

ws4.cell(row=1, column=1, value='精算险类代码')
ws4.cell(row=1, column=2, value='精算险类名称')
for c in range(1, 3):
    cell = ws4.cell(row=1, column=c)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = thin_border

for r_idx, (code, name) in enumerate(ACTUARIAL_CLASSES, 2):
    ws4.cell(row=r_idx, column=1, value=code)
    ws4.cell(row=r_idx, column=2, value=name)
    for c in range(1, 3):
        ws4.cell(row=r_idx, column=c).border = thin_border
        ws4.cell(row=r_idx, column=c).font = normal_font

ws4.column_dimensions['A'].width = 16
ws4.column_dimensions['B'].width = 24

wb.save(OUT_FILE)
print(f'已生成模板: {OUT_FILE}')
