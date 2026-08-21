#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
从【当前附件】 IFRS17财务预测_改造版_开发.xlsx 直接解析六张输出工作表，
重建权威目标 fixture：data_input/fixtures/verify_targets.json

用途：让系统「计量结果输出」与「验证模块」六表直接以附件（当前schema）为目标，
字段名/列序/数值与附件 1:1 对齐（旧 verify_targets.json 来自 0729_VBA 旧附件，
字段还是 合同组ID 等旧 schema，与当前附件 预测组ID 不一致）。

解析逻辑与 data_input/views.py 的 api_upload_verify 完全一致：
- 普通表用 _parse_verify_sheet（第1行表头，截尾 None，跳过全空行）
- 财务报表类用 _parse_financial_report_sheet（前2行标题，第4行日期表头，第5行起数据，科目在B列）
"""
import json
import os

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
FIXTURE_DIR = os.path.join(BASE_DIR, 'ifrs17-system', 'data_input', 'fixtures')
OUTPUT_FILE = os.path.join(FIXTURE_DIR, 'verify_targets.json')

# 默认附件路径（可用命令行参数覆盖）
DEFAULT_ATT = os.path.join(BASE_DIR, '验证文件', 'IFRS17财务预测_改造版_开发.xlsx')

VERIFY_OUTPUT_SHEETS = [
    'PAA计算_新业务整理',
    'PAA计算_现有业务整理',
    'PAA计算_汇总',
    'PAA计算_MTD',
    '输出财务报表_MTD',
    '输出财务报表_YTD',
]
FINANCIAL_REPORT_SHEETS = ('输出财务报表', '输出财务报表_MTD', '输出财务报表_YTD')


def _norm_header(h):
    """表头键统一为字符串；日期转为 YYYY-MM-DD。"""
    if h is None:
        return None
    if hasattr(h, 'strftime'):
        return h.strftime('%Y-%m-%d')
    return str(h)


def _parse_verify_sheet(ws):
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_row = next(rows_iter)
    except StopIteration:
        return [], []
    headers = [_norm_header(h) for h in header_row]
    last_non_none = 0
    for i, h in enumerate(headers):
        if h is not None:
            last_non_none = i + 1
    headers = headers[:last_non_none]
    data_rows = []
    for row in rows_iter:
        if all(v is None for v in row):
            continue
        row_list = list(row)[:len(headers)]
        while len(row_list) < len(headers):
            row_list.append(None)
        data_rows.append(row_list)
    return headers, data_rows


def _parse_financial_report_sheet(ws):
    all_rows = [list(r) for r in ws.iter_rows(values_only=True)]
    if len(all_rows) < 4:
        return [], []
    date_row = all_rows[3] if len(all_rows) > 3 else []
    headers = ['科目']
    date_idx = []
    for i in range(2, len(date_row)):
        v = date_row[i]
        if v is None:
            continue
        headers.append(v.strftime('%Y-%m-%d') if hasattr(v, 'strftime') else str(v).strip())
        date_idx.append(i)
    data_rows = []
    for row in all_rows[4:]:
        if all(v is None for v in row):
            continue
        row_list = [None] * len(headers)
        row_list[0] = row[1] if len(row) > 1 else None
        for j, di in enumerate(date_idx):
            row_list[1 + j] = row[di] if di < len(row) else None
        data_rows.append(row_list)
    return headers, data_rows


def _rows_to_dicts(headers, rows):
    out = []
    for r in rows:
        out.append({h: (r[i] if i < len(r) else None) for i, h in enumerate(headers)})
    return out


def main():
    import sys
    att = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_ATT
    if not os.path.exists(att):
        print(f'[error] 附件不存在: {att}')
        sys.exit(1)

    from openpyxl import load_workbook
    wb = load_workbook(att, read_only=True, data_only=True)

    targets = {}
    meta = {}
    missing = []
    for sn in VERIFY_OUTPUT_SHEETS:
        if sn not in wb.sheetnames:
            missing.append(sn)
            print(f'[warn] 附件缺少工作表: {sn}')
            continue
        ws = wb[sn]
        if sn in FINANCIAL_REPORT_SHEETS:
            headers, rows = _parse_financial_report_sheet(ws)
        else:
            headers, rows = _parse_verify_sheet(ws)
        targets[sn] = _rows_to_dicts(headers, rows)
        meta[sn] = {'rows': len(rows), 'cols': len(headers), 'headers': headers}
        print(f'[ok] {sn}: {len(rows)} rows x {len(headers)} cols')

    wb.close()

    if missing:
        print(f'[warn] 缺少的工作表将不会出现在 fixture 中: {missing}')

    os.makedirs(FIXTURE_DIR, exist_ok=True)
    with open(OUTPUT_FILE, 'w', encoding='utf-8') as f:
        json.dump({'sheets': targets, 'meta': meta, 'source': os.path.basename(att)},
                  f, ensure_ascii=False, indent=None, default=str)
    print(f'[done] 写入 {OUTPUT_FILE}')


if __name__ == '__main__':
    main()
