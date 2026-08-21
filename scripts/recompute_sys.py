# -*- coding: utf-8 -*-
import os, sys, django
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'ifrs17_core.settings')
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
django.setup()
from data_input.models import SheetData
import openpyxl

ROOT="D:/IFRS17预测模型"
VP=os.path.join(ROOT,"验证文件","IFRS17财务预测_改造版_0729_VBA.xlsx")
wb=openpyxl.load_workbook(VP,data_only=True)

def srow(sheet, key, idx):
    s=SheetData.objects.filter(is_active=True, sheet_name=sheet).first()
    for r in s.rows:
        if idx<len(r) and r[idx] is not None and str(r[idx])==str(key):
            return {h:r[i] for i,h in enumerate(s.headers) if h is not None}
    return {}

def vrow(sheet, key, idx):
    ws=wb[sheet]; rows=list(ws.iter_rows(values_only=True)); hdr=rows[0]
    for r in rows[1:]:
        if idx<len(r) and r[idx] is not None and str(r[idx])==str(key):
            return {h:r[i] for i,h in enumerate(hdr) if h is not None}
    return {}

def mdict(d):
    return {int(float(k)): float(v) for k,v in d.items() if k is not None and str(k).replace('.','',1).isdigit()}

g='交强险_直保或分入_现有业务'
gid='交强险_直保或分入_现有业务&0'
ac=32
# system inputs
sys_op=srow('系统期初余额表', g, 2)
sys_ep=srow('预期赔付率', g, 2)
sys_earn=srow('未到期赚取模式_现有业务', g, 3)
sys_und=srow('未到期赔付模式', ac, 3)
sys_dev=srow('未决赔付模式', ac, 3)
sys_cum=srow('实际赔付比例', ac, 3)
# verification inputs
v_op=vrow('系统期初余额表', g, 2)
v_ep=vrow('输入整理-预期赔付率', g, 2)
v_earn=vrow('输入整理-未到期赚取模式', gid, 3)
v_und=vrow('输入整理-未到期赔付模式', ac, 3)
v_dev=vrow('输入整理-未决赔付模式', ac, 3)
v_cum=vrow('输入整理-实际赔付比例_累计', ac, 3)

print("open_X  sys/ver:", sys_op.get('已发生未决赔款负债_预期现金流'), v_op.get('已发生未决赔款负债_预期现金流'))
print("pr p1    sys/ver:", sys_ep.get('1'), v_ep.get('1'))
print("earn     sys:", mdict(sys_earn), "\n         ver:", mdict(v_earn))
print("und_incur sys:", mdict(sys_und), "\n         ver:", mdict(v_und))
print("dev      sys:", mdict(sys_dev), "\n         ver:", mdict(v_dev))
print("cum      sys:", mdict(sys_cum), "\n         ver:", mdict(v_cum))

# verification reported AC
wz=wb["PAA计算_现有业务整理"]; rows=list(wz.iter_rows(values_only=True)); hdr=rows[0]
for r in rows[1:]:
    if r[3]==gid:
        print("VERIFY AC_期末预期现金流:", r[hdr.index('输出_已发生未决赔款负债_预期现金流')])
        break
wb.close()
