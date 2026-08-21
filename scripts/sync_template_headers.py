"""Sync affected template sheet headers to latest excel_specs definitions.

Uses data_input/excel_specs.py as the single source of truth so the upload
templates always match what excel_validator.py enforces.
"""
import os
import sys
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, 'ifrs17-system'))
from data_input.excel_specs import SYSTEM_DOCK_SPECS, EXCEL_UPLOAD_SPECS

SYSTEM_DOCK = os.path.join(ROOT, '验证文件', 'backup', '系统对接输入表_模板.xlsx')
EXCEL_UPLOAD = os.path.join(ROOT, '验证文件', 'backup', '手工输入表_模板.xlsx')


def sync_sheet(wb, spec_map, sheet_name, numeric_fill):
    """Replace row 1 of `sheet_name` with spec headers (+ numeric cols if numeric_fill)."""
    if sheet_name not in spec_map:
        print(f"  [SKIP] '{sheet_name}' 不在规范中")
        return
    spec = spec_map[sheet_name]
    text = list(spec.headers)
    if sheet_name not in wb.sheetnames:
        print(f"  [SKIP] '{sheet_name}' 不在模板中")
        return
    ws = wb[sheet_name]
    total = len(text) + (numeric_fill if numeric_fill else 0)
    for c in range(1, total + 1):
        val = text[c - 1] if c <= len(text) else (c - len(text))
        ws.cell(row=1, column=c).value = val
    # 清除超出 total 的残留表头（如旧模板多出的 减值 / 合同组合名称 等）
    max_col = ws.max_column
    for c in range(total + 1, max_col + 1):
        ws.cell(row=1, column=c).value = None
    print(f"  [OK] '{sheet_name}'  row1({total}列)={text}" + (f" +1..{numeric_fill}" if numeric_fill else ""))


print(">>> 系统对接输入表_模板.xlsx")
wb1 = openpyxl.load_workbook(SYSTEM_DOCK)
# 系统期初余额表: 纯文本表头（无数值列），按规范 45 字段
sync_sheet(wb1, SYSTEM_DOCK_SPECS, '系统期初余额表', None)
# 三张现有业务模式表模板已用 预测组，且文本列与规范一致，无需改
wb1.save(SYSTEM_DOCK)
wb1.close()

print(">>> 手工输入表_模板.xlsx")
wb2 = openpyxl.load_workbook(EXCEL_UPLOAD)
sync_sheet(wb2, EXCEL_UPLOAD_SPECS, '对应关系配置表', None)
sync_sheet(wb2, EXCEL_UPLOAD_SPECS, '合同组拼接', None)
wb2.save(EXCEL_UPLOAD)
wb2.close()
print(">>> 完成")
