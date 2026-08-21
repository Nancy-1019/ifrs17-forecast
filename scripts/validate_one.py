# -*- coding: utf-8 -*-
import os, datetime, calendar
from openpyxl import load_workbook
ROOT="D:/IFRS17预测模型"
VBX=os.path.join(ROOT,"验证文件","IFRS17财务预测_改造版_0729_VBA.xlsx")
wb=load_workbook(VBX,data_only=True)

# verification computed values for saved group
ws=wb["现有业务预期现金流_1"]
rows=list(ws.iter_rows(values_only=True))
hdr=rows[0]
def ci(c):
    n=0
    for ch in c: n=n*26+(ord(ch)-64)
    return n-1
# column indices by header
def col(name): return hdr.index(name)
# print first 65 rows' AW (预期保费现金流), AI(事故日_维持), a(AD), d(AE)
print("=== saved triangle: a(AD),d(AE),AW(预期保费现金流),AI(事故日_维持) first 65 rows ===")
for ri in range(1,66):
    r=rows[ri]
    a=r[col("事故期月度序列")]; d=r[col("发展期月度序列")]
    aw=r[col("预期保费现金流")]; ai=r[col("事故日_维持")]
    print(f" row{ri+1}: a={a} d={d} AW={aw} AI={ai}")

# 预测表 row2 (other_分出): AO, AU, BI
wf=wb["PAA计算_现有业务预测"]
rf=list(wf.iter_rows(values_only=True))
hf=rf[0]
def fcol(n): return hf.index(n)
row2=rf[1]
print("\n=== 预测表 row2 (其他_分出_现有业务&0) key cols ===")
print(" 当期收到的保费(AO)=", row2[fcol("当期收到的保费")])
print(" 期末未到期责任负债_非亏损部分(AU)=", row2[fcol("期末未到期责任负债_非亏损部分")])
print(" 输出_未到期责任负债_非亏损部分(BI)=", row2[fcol("输出_未到期责任负债_非亏损部分")])
print(" 输出_已发生未决赔款负债_预期现金流(BL)=", row2[fcol("输出_已发生未决赔款负债_预期现金流")])
print(" 输出_现金流_收到的保费(CJ)=", row2[fcol("输出_现金流_收到的保费")])

# 整理 for this group
wz=wb["PAA计算_现有业务整理"]; rz=list(wz.iter_rows(values_only=True)); hz=rz[0]
for r in rz[1:]:
    if r[3]=="其他_分出_现有业务&0" and r[4]=="分出":
        print("\n=== 整理 row for 其他_分出_现有业务&0 ===")
        print(" 输出_未到期责任负债_非亏损部分=", r[hz.index("输出_未到期责任负债_非亏损部分")])
        print(" 输出_现金流_收到的保费=", r[hz.index("输出_现金流_收到的保费")])
        print(" 输出_已发生未决赔款负债_预期现金流=", r[hz.index("输出_已发生未决赔款负债_预期现金流")])
        break
wb.close()
