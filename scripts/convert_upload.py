"""将 IFRS17财务预测_改造版_0729_VBA.xlsx 中的数据整理为系统上传文件。

核心原则（对齐 oracle vba_faithful.py）：
  - 现有业务（及共享）模式/比率表，VBA 实际引用的是「输入整理-*」加工表，
    而非同名原始表。加工表列头存在错位/别名（如 维持费用率 的 险类代码 写在“合同组ID”列），
    因此本脚本按“规范表头 -> 加工表源列下标”显式映射，再输出规范表头。
  - 每个 预测组 在加工模式表中恰对应 1 个 合同组ID（已验证），故系统按 预测组/精算险类
    取数等价于 oracle 按 合同组ID 取数。
  - 加工表同时含 _2026(新业务) 与 _现有业务 行，故新业务同名模式表也改从加工表取数（数据正确，
    新业务引擎逻辑另需移植）。
  - 保留从原始表读取的表：系统期初余额表 / 预期摊回比例_现有业务 / 财务报表实际数 /
    初始确认利率曲线 / 即期利率曲线 / 风险调整比例 / 再保人不履约风险 / 投资成分比例 / 新业务专属表等。
"""
import os
import sys
import datetime

import openpyxl
from openpyxl import Workbook

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PROJECT_ROOT = os.path.dirname(ROOT)
sys.path.insert(0, ROOT)
from data_input.excel_specs import EXCEL_UPLOAD_SPECS, SYSTEM_DOCK_SPECS  # noqa: E402
from data_input.excel_validator import validate_upload  # noqa: E402

SRC = os.path.join(PROJECT_ROOT, "验证文件", "IFRS17财务预测_改造版_0729_VBA.xlsx")
OUT_EXCEL = os.path.join(PROJECT_ROOT, "验证文件", "手工输入表_上传数据_0729.xlsx")
OUT_DOCK = os.path.join(PROJECT_ROOT, "验证文件", "系统对接输入表_上传数据_0729.xlsx")

# 加工表 -> 规范表头 的源列映射（None 表示该规范列在加工表中无对应，留空）。
# 未列出的规范表头若能在加工表按名匹配则按名取，否则留空。
# numeric 默认取加工表“整数月”列头（前 60 个）；可用 numeric_idxs 显式指定。
PROCESS_MAP = {
    # ---- 系统对接（现有业务） ----
    '合同组关键假设_现有业务': dict(
        src='输入整理-关键假设',
        col_map={'更新日期': 0, '评估时点': 1, '合同组ID': 4, '精算险类': 7,
                 '评估期开始业务预期写入时间（月）': 10, '子合同组合名称': 5,
                 '业务类型': 8, '应收保费': 14, '应付IACF': 15},
    ),
    '保费现金流模式_现有业务': dict(src='输入整理-保费现金流模式', col_map={'合同组ID': 3}),
    'IACF现金流模式_现有业务': dict(src='输入整理-IACF现金流模式', col_map={'合同组ID': 3, '预测组': 4}),
    '未到期赚取模式_现有业务': dict(src='输入整理-未到期赚取模式', col_map={'合同组ID': 3, '预测组': 4}),
    '提前确认保单_现有业务': dict(
        src='提前确认保单_现有业务',
        col_map={'更新日期': 0, '评估时点': 1, '数据类型': 2, '精算险类': 3},
        numeric_idxs=list(range(4, 64)),  # 60 个日期列头 -> 重命名为 1..60
    ),
    # ---- 手工输入（共享/新业务模式表） ----
    '保费现金流模式_新业务': dict(src='输入整理-保费现金流模式'),
    'IACF现金流模式_新业务': dict(src='输入整理-IACF现金流模式'),
    '未到期赚取模式_新业务': dict(src='输入整理-未到期赚取模式'),
    '预期赔付率': dict(src='输入整理-预期赔付率'),
    '维持费用率': dict(src='输入整理-维持费用率', col_map={'精算险类': 4}),  # 源idx4=险类代码(标为合同组ID)
    '未到期间接理赔费用率': dict(src='输入整理-未到期间接理赔费用率', col_map={'精算险类': 4}),
    '未决间接理赔费用率': dict(src='输入整理-未决间接理赔费用率', col_map={'精算险类': 4}),
    '未到期赔付模式': dict(src='输入整理-未到期赔付模式'),
    '未决赔付模式': dict(src='输入整理-未决赔付模式'),
    '实际赔付比例': dict(src='输入整理-实际赔付比例_累计'),
}


def _clean(v):
    if v is None:
        return None
    if isinstance(v, str):
        s = v.strip()
        return s if s else None
    return v


def _read_sheet_rows(wb, sheet_name):
    ws = wb[sheet_name]
    rows = []
    for r in ws.iter_rows(values_only=True):
        rows.append(list(r))
    return rows


def _is_int_month(h):
    if h is None:
        return False
    try:
        v = float(str(h).strip())
        return v == int(v) and v > 0
    except (ValueError, TypeError):
        return False


def build_sheet(src_rows, src_header, spec, col_map=None, numeric_idxs=None):
    """从加工表(src)按规范(spec)抽取：规范表头 + 60 月数值。"""
    col_map = col_map or {}
    src_name = {str(h).strip(): i for i, h in enumerate(src_header) if h is not None}
    resolved = {}
    for h in spec.headers:
        if h in col_map:
            resolved[h] = col_map[h]
        elif h in src_name:
            resolved[h] = src_name[h]
        else:
            resolved[h] = None
    if numeric_idxs is None:
        numeric_idxs = [i for i, h in enumerate(src_header) if _is_int_month(h)]
        if spec.numeric_cols:
            numeric_idxs = numeric_idxs[:spec.numeric_cols]
    out_header = list(spec.headers) + [str(i + 1) for i in range(len(numeric_idxs))]
    out_rows = []
    for r in src_rows[1:]:
        out = []
        for h in spec.headers:
            si = resolved[h]
            out.append(_clean(r[si]) if (si is not None and si < len(r)) else None)
        for ni in numeric_idxs:
            out.append(_clean(r[ni]) if ni < len(r) else None)
        out_rows.append(out)
    return out_header, out_rows


def convert_contract_sheet(rows, spec):
    """初始确认利率曲线：第 3 行为真实表头（合同组名为后续列）。"""
    real_hdr = rows[2]
    while real_hdr and (real_hdr[-1] is None or str(real_hdr[-1]).strip() == ""):
        real_hdr = real_hdr[:-1]
    out_header = list(spec.headers)
    group_idxs = [i for i in range(len(spec.headers), len(real_hdr))
                  if isinstance(real_hdr[i], str) and str(real_hdr[i]).strip() != ""]
    out_header += [real_hdr[i] for i in group_idxs]
    out_rows = []
    for dr in rows[3:]:
        dr = list(dr)
        width = len(out_header)
        if len(dr) < width:
            dr = dr + [None] * (width - len(dr))
        else:
            dr = dr[:width]
        out_rows.append(dr)
    return out_header, out_rows


def build_workbook(specs, source, out_path):
    print(f"\n=== {source}: 生成 {out_path} ===")
    wb_src = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    wb_out = Workbook()
    wb_out.remove(wb_out.active)

    all_ok = True
    for key, spec in specs.items():
        cfg = PROCESS_MAP.get(key)
        if cfg is not None:
            src_sheet = cfg['src']
        else:
            src_sheet = key
        if src_sheet not in wb_src.sheetnames:
            print(f"  [缺失] {src_sheet} 不在附件中 (目标表 {key})")
            all_ok = False
            continue
        rows = _read_sheet_rows(wb_src, src_sheet)
        src_header = rows[0]
        while src_header and (src_header[-1] is None or str(src_header[-1]).strip() == ""):
            src_header = src_header[:-1]

        if spec.contract_cols:
            out_headers, out_rows = convert_contract_sheet(rows, spec)
        else:
            col_map = cfg.get('col_map') if cfg else None
            numeric_idxs = cfg.get('numeric_idxs') if cfg else None
            out_headers, out_rows = build_sheet(rows, src_header, spec, col_map, numeric_idxs)

        ws = wb_out.create_sheet(title=key)
        ws.append(out_headers)
        for r in out_rows:
            ws.append(r)
        # 基本自检
        warn = []
        if spec.numeric_cols and len(out_headers) < len(spec.headers) + spec.numeric_cols:
            warn.append(f"列数不足: {len(out_headers)} < {len(spec.headers)+spec.numeric_cols}")
        if spec.sum_to_one:
            for ri, r in enumerate(out_rows):
                s = 0.0
                cnt = 0
                for ci in range(len(spec.headers), len(spec.headers) + spec.numeric_cols):
                    if ci < len(r) and isinstance(r[ci], (int, float)):
                        s += float(r[ci]); cnt += 1
                if cnt and abs(s - 1.0) > 0.0001:
                    warn.append(f"第{ri+2}行 模式和不等于1: {s:.6f}")
                    break
        flag = "OK" if not warn else "WARN"
        print(f"  [{flag}] {key}: {len(out_rows)} 行, {len(out_headers)} 列"
              + (f" | {'; '.join(warn)}" if warn else ""))
        if warn:
            all_ok = False

    wb_src.close()
    wb_out.save(out_path)
    print(f"  保存完成: {os.path.getsize(out_path)} 字节")
    return all_ok


def validate_file(path, source):
    print(f"\n--- 校验 {path} (source={source}) ---")
    with open(path, "rb") as f:
        data = f.read()
    res = validate_upload(data, source)
    print(f"  success={res['success']}")
    if res["errors"]:
        print(f"  errors: {res['errors']}")
    if res["warnings"]:
        print(f"  warnings: {res['warnings']}")
    bad = [(k, v) for k, v in res["sheetResults"].items() if not v.get("valid")]
    if bad:
        for k, v in bad:
            print(f"  [INVALID] {k}: {v.get('errors')}")
    else:
        print(f"  全部 {len(res['sheetResults'])} 张表通过校验")
    return res["success"]


def main():
    ok1 = build_workbook(EXCEL_UPLOAD_SPECS, "excel", OUT_EXCEL)
    ok2 = build_workbook(SYSTEM_DOCK_SPECS, "dock", OUT_DOCK)
    v1 = validate_file(OUT_EXCEL, "excel")
    v2 = validate_file(OUT_DOCK, "dock")
    print("\n========== 汇总 ==========")
    print(f"手工输入表 生成{'成功' if ok1 else '有告警'} / 校验{'通过' if v1 else '失败'}")
    print(f"系统对接表 生成{'成功' if ok2 else '有告警'} / 校验{'通过' if v2 else '失败'}")


if __name__ == "__main__":
    main()
