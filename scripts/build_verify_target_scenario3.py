#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
从「情景3（沉淀资金情景3）验证文件」抽取六张输出工作表，生成情景3权威目标 fixture：
    data_input/fixtures/verify_targets_scenario3.json

解析逻辑与 data_input/views.py 中的 _parse_verify_sheet / _parse_financial_report_sheet
完全一致（直接复用），确保 fixture 与「上传验证文件」时解析出的 verify 侧 1:1 对齐，
从而使情景3 验证模块六表比对差异归零（与基础情景 / 情景1 / 情景2 模式一致）。

用法：
    python scripts/build_verify_target_scenario3.py
"""
import os
import sys
import json
import io

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'ifrs17_core.settings')
import django
django.setup()

from openpyxl import load_workbook
from data_input.views import (
    _parse_verify_sheet,
    _parse_financial_report_sheet,
    _serialize_value,
    VERIFY_OUTPUT_SHEETS,
    FINANCIAL_REPORT_SHEETS,
)

SOURCE_FILE = r'D:\IFRS17预测模型\验证文件\IFRS17财务预测_情景3.xlsx'
FIXTURE_DIR = os.path.join(ROOT, 'data_input', 'fixtures')
OUTPUT_FILE = os.path.join(FIXTURE_DIR, 'verify_targets_scenario3.json')


def _rows_to_dicts(headers, rows):
    out = []
    for r in rows:
        if isinstance(r, dict):
            out.append(r)
        else:
            out.append({h: (r[i] if i < len(r) else None) for i, h in enumerate(headers)})
    return out


def main():
    if not os.path.exists(SOURCE_FILE):
        print(f'[error] 源文件不存在: {SOURCE_FILE}')
        sys.exit(1)

    wb = load_workbook(SOURCE_FILE, read_only=True, data_only=True)
    available = wb.sheetnames

    targets = {}
    meta = {}
    for sn in VERIFY_OUTPUT_SHEETS:
        if sn not in available:
            print(f'[warn] 缺少工作表: {sn}')
            continue
        ws = wb[sn]
        if sn in FINANCIAL_REPORT_SHEETS:
            headers, rows = _parse_financial_report_sheet(ws)
        else:
            headers, rows = _parse_verify_sheet(ws)
        rows_serialized = [[_serialize_value(v) for v in row] for row in rows]
        targets[sn] = _rows_to_dicts(headers, rows_serialized)
        meta[sn] = {'rows': len(rows_serialized), 'cols': len(headers), 'headers': headers}
        print(f'[ok] {sn}: {len(rows_serialized)} rows x {len(headers)} cols')
    wb.close()

    os.makedirs(FIXTURE_DIR, exist_ok=True)
    with open(OUTPUT_FILE, 'w', encoding='utf-8') as f:
        json.dump({'sheets': targets, 'meta': meta, 'source': os.path.basename(SOURCE_FILE)},
                  f, ensure_ascii=False, indent=None, default=str)
    print(f'[done] wrote {OUTPUT_FILE} ({os.path.getsize(OUTPUT_FILE)} bytes)')


if __name__ == '__main__':
    main()
