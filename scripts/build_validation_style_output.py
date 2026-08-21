#!/usr/bin/env python
"""
按验证文件输出格式（PAA计算_新业务整理 / PAA计算_现有业务整理）
重建系统计量输出样式，生成 Excel：
  D:/IFRS17预测模型/验证文件/系统计量输出_验证格式.xlsx

规则：
- 列名、列序与验证文件完全一致；
- 系统已有的字段直接映射；
- 系统没有的中间字段用 0 占位（新业务期初现值、累计现金流等第一期天然为 0）；
- 输出带表头冻结、千分位数字格式。
"""
import os
import json
import datetime
from copy import deepcopy
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
CACHE_PATH = os.path.join(ROOT, 'ifrs17-system', 'data_input', 'runtime_cache', 'calc_result_cache.json')
VERIFY_FILE = os.path.join(ROOT, '验证文件', 'IFRS17财务预测_改造版_0729_VBA.xlsx')
OUT_FILE = os.path.join(ROOT, '验证文件', '系统计量输出_验证格式.xlsx')

# 尝试读取本地 openpyxl
from openpyxl import load_workbook


def _serialize(v):
    if v is None:
        return None
    if isinstance(v, datetime.datetime):
        return v.strftime('%Y-%m-%d %H:%M:%S')
    if isinstance(v, datetime.date):
        return v.strftime('%Y-%m-%d')
    return v


def _to_float(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(v)
    except Exception:
        return None


def load_cache():
    with open(CACHE_PATH, encoding='utf-8') as f:
        return json.load(f).get('result', {})


def get_validation_headers(sheet_name):
    wb = load_workbook(VERIFY_FILE, read_only=True, data_only=True)
    if sheet_name not in wb.sheetnames:
        wb.close()
        return None
    ws = wb[sheet_name]
    hdr = [str(h) if h is not None else '' for h in next(ws.iter_rows(values_only=True))]
    wb.close()
    return hdr


def base_group_name(pid):
    """交强险_直保或分入_2026&1 -> 交强险_直保或分入_2026"""
    if not pid:
        return ''
    pid = str(pid)
    if '&' in pid:
        return pid.rsplit('&', 1)[0]
    return pid


# 新业务字段映射：{验证列名: 系统字段名或 callable(row)}
NB_FIELD_MAP = {
    '子合同组合名称': lambda r: str(r.get('精算险类', '')),
    '合同组合名称': lambda r: 1,
    '合同组名称': lambda r: base_group_name(r.get('预测组ID', '')),
    '合同组ID': lambda r: r.get('预测组ID', ''),
    '直保或分入/分出': '直保或分入/分出',
    '精算险类': '精算险类',
    '评估日': '评估日',
    '预测时点': '预测时点',
    '预测时点年初': '预测时点年初',
    '预测间隔': '预测间隔',
    '数据来源表': '数据来源表',
    '当期确认比例': '当期确认比例',
    '获取现金流摊销比例': '获取现金流摊销比例',
    '投资成分比例%': '投资成分比例%',
    '预期摊回比例%': '预期摊回比例%',
    '当月计息利率': '当月计息利率',
    '保费收入': '保费收入',
    '获取费用': '获取费用',
    # 现金流与输出列
    '当期收到的保费': '输出_现金流_收到的保费',
    '当期支出的保险获取现金流': '输出_现金流_支付的IACF',
    '未到期责任负债计息': '输出_IFIE_未到期_未到期计息',
    '当期确认的保险收入': lambda r: -_to_float(r.get('输出_保险合同收入', 0)) if r.get('输出_保险合同收入') is not None else 0,
    '摊销的获取费用': '输出_赔付与费用_摊销的保险获取现金流',
    '当期分解的投资成分': '输出_赔付与费用_分解的投资成分',
    '期末未到期责任负债_非亏损部分': '输出_未到期责任负债_非亏损部分',
    '未到期责任负债_亏损部分': '输出_未到期责任负债_亏损部分',
    '未到期责任负债_亏损摊回': '输出_未到期责任负债_亏损摊回',
    '现金流_支付的维持费用': '输出_现金流_支付的维持费用',
    '现金流_支付的赔付与理赔费用': '输出_现金流_支付的赔付与理赔费用',
}

# 现有业务字段映射
EB_FIELD_MAP = {
    '子合同组合名称': lambda r: str(r.get('精算险类', '')),
    '合同组合名称': lambda r: 1,
    '合同组名称': lambda r: base_group_name(r.get('合同组ID', '')),
    '合同组ID': '合同组ID',
    '现有业务预测组': '现有业务预测组',
    '现有业务预测组ID': '现有业务预测组ID',
    '预测组': '预测组',
    '预测组ID': '预测组ID',
    '直保或分入/分出': '直保或分入/分出',
    '精算险类': '精算险类',
    '评估日': '评估日',
    '预测时点': '预测时点',
    '预测时点年初': '预测时点年初',
    '预测间隔': '预测间隔',
    '数据来源表': '数据来源表',
    '当期确认比例': '当期确认比例',
    '获取现金流摊销比例': '获取现金流摊销比例',
    '当月计息利率': '当月计息利率',
    # 现值/期末
    '期末_预期现金流_已发生未决赔款负债': '期末现值_未决赔款现金流_期末',
    '期末_预期现金流_间接理赔费用负债': '期末现值_未决间接理赔费用现金流_期末',
    '期末_预期现金流_已发生未决赔款负债_再保人不履约': '期末现值_未决赔款现金流_再保人不履约风险_期末',
    '期末_非金融风险调整_已发生未决赔款负债': '期末现值_未决赔款现金流_非金融风险调整_期末',
    '期末_非金融风险调整_间接理赔费用负债': '期末现值_未决间接理赔费用现金流_非金融风险调整_期末',
    '期末_非金融风险调整_已发生未决赔款负债_再保人不履约': '期末现值_未决赔款现金流_再保人不履约风险_非金融风险调整_期末',
    # 现金流与输出列
    '当期收到的保费': '输出_现金流_收到的保费',
    '当期支出的保险获取现金流': '输出_现金流_支付的IACF',
    '未到期责任负债计息': '输出_IFIE_未到期_未到期计息',
    '当期确认的保险收入': lambda r: -_to_float(r.get('输出_保险合同收入', 0)) if r.get('输出_保险合同收入') is not None else 0,
    '摊销的获取费用': '输出_赔付与费用_摊销的保险获取现金流',
    '当期分解的投资成分': '输出_赔付与费用_分解的投资成分',
    '期末未到期责任负债_非亏损部分': '输出_未到期责任负债_非亏损部分',
    '未到期责任负债_亏损部分': '输出_未到期责任负债_亏损部分',
    '未到期责任负债_亏损摊回': '输出_未到期责任负债_亏损摊回',
    '现金流_支付的赔付': '现金流_支付的赔付',
    '现金流_支付的理赔费用': '现金流_支付的理赔费用',
    '现金流_支付的赔付_累计': '现金流_支付的赔付_累计',
    '现金流_支付的间接理赔费用_累计': '现金流_支付的间接理赔费用_累计',
    '现金流_支付的维持费用': '输出_现金流_支付的维持费用',
}

# 输出_* 列直接同名映射，两边通用
OUTPUT_PREFIX = '输出_'


def build_styled_rows(sys_rows, hdr, field_map):
    out = []
    for r in sys_rows:
        new = {}
        for col in hdr:
            val = None
            if col.startswith(OUTPUT_PREFIX):
                val = r.get(col)
            elif col in field_map:
                m = field_map[col]
                if callable(m):
                    try:
                        val = m(r)
                    except Exception:
                        val = None
                else:
                    val = r.get(m)
            else:
                # 未映射字段默认 0（保持与验证文件同列数）
                val = 0
            new[col] = val
        out.append(new)
    return out


def write_excel(sheets_data):
    wb = Workbook()
    wb.remove(wb.active)
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill('solid', fgColor='4472C4')
    header_align = Alignment(horizontal='center', vertical='center', wrap_text=True)
    thin = Side(style='thin', color='CCCCCC')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    num_fmt = '#,##0.00'

    for sheet_name, rows, hdr in sheets_data:
        ws = wb.create_sheet(title=sheet_name)
        ws.append(hdr)
        for cell in ws[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
            cell.border = border
        ws.freeze_panes = 'A2'

        for r in rows:
            ws.append([_serialize(r.get(c)) for c in hdr])

        # 简单列宽与数字格式
        for i, col in enumerate(hdr, 1):
            letter = get_column_letter(i)
            max_len = max(len(str(col)), 12)
            ws.column_dimensions[letter].width = min(max_len + 2, 30)
            # 数值列统一两位小数
            for cell in ws[letter][1:]:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = num_fmt
                cell.border = border

    wb.save(OUT_FILE)
    print(f"已生成验证格式系统输出: {OUT_FILE}")


def main():
    cache = load_cache()
    sheets = [
        ('PAA计算_新业务整理', cache.get('newBusinessPredict', []), NB_FIELD_MAP),
        ('PAA计算_现有业务整理', cache.get('existingBusinessPredict', []), EB_FIELD_MAP),
    ]
    data = []
    for vname, sys_rows, fmap in sheets:
        hdr = get_validation_headers(vname)
        if hdr is None:
            print(f"验证文件中缺少 {vname}，跳过")
            continue
        if not sys_rows:
            print(f"系统缓存中缺少 {vname} 数据，跳过")
            continue
        rows = build_styled_rows(sys_rows, hdr, fmap)
        data.append((vname, rows, hdr))
        print(f"{vname}: 系统行数={len(sys_rows)} -> 输出列数={len(hdr)}")
    if data:
        write_excel(data)
    else:
        print("无数据可输出")


if __name__ == '__main__':
    main()
