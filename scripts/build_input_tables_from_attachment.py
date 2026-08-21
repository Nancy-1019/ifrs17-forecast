# -*- coding: utf-8 -*-
"""
根据系统 手工输入表 / 系统对接输入表 模板结构，
把附件 IFRS17财务预测_改造版_开发.xlsx 中的数据抽取并整理成两版输入表：

  验证文件/手工输入表_附件数据版.xlsx
  验证文件/系统对接输入表_附件数据版.xlsx

规则：
  - 输出列顺序/列名 = 系统模板（即 excel_specs.py 权威定义）。
  - 数据按列名从附件映射；附件列名与模板一致的直接填充。
  - 费用输入项 / 其他输入项：附件月度列是“月末日期”，模板是序数 1..60，
    按位置（同序）映射，输出列名改为 1..60。
  - 系统期初余额表 / 对应关系配置表 / 合同组拼接：模板多出的列在附件中缺失，置空。
  - 初始确认利率曲线：附件为 3 行表头，第 3 行才是规范表头（含合同组列），
    以第 3 行为表头、第 4 行起为数据。
  - 说明 表：取附件“输入表说明”内容。
"""
import openpyxl
import datetime

ATT = "验证文件/IFRS17财务预测_改造版_开发.xlsx"
DOCK_T = "验证文件/backup/系统对接输入表_模板.xlsx"
MAN_T = "验证文件/backup/手工输入表_模板.xlsx"
OUT_DOCK = "验证文件/系统对接输入表_附件数据版.xlsx"
OUT_MAN = "验证文件/手工输入表_附件数据版.xlsx"

# 模板 sheet -> 附件 sheet 名称映射（默认同名）
MANUAL_NAME_MAP = {"说明": "输入表说明"}
DOCK_NAME_MAP = {"说明": "输入表说明"}

# 附件表头所在行（1-indexed）。初始确认利率曲线为第3行。
ATT_HEADER_ROW = {"初始确认利率曲线": 3}

# 月度列按位置映射（日期 -> 序数）的表
DATE_TO_ORDINAL_SHEETS = {"费用输入项", "其他输入项"}


def read_headers(ws, row=1):
    out = []
    for c in range(1, ws.max_column + 1):
        out.append(ws.cell(row=row, column=c).value)
    while out and out[-1] is None:
        out.pop()
    return out


def trim_header(hdr, sheet_name):
    """初始确认利率曲线第3行表头后存在 padding 的 0（非列名），
    截取到最后一个字符串列名处。"""
    if sheet_name == "初始确认利率曲线":
        k = len(hdr)
        while k > 0 and not isinstance(hdr[k - 1], str):
            k -= 1
        return hdr[:k]
    return hdr


def load_attachment_sheets(sheet_list):
    """返回 {sn: (headers, [rows...])} 。headers 取对应表头行，rows 为数据行。"""
    wb = openpyxl.load_workbook(ATT, read_only=True, data_only=True)
    res = {}
    for sn in sheet_list:
        if sn not in wb.sheetnames:
            continue
        ws = wb[sn]
        hrow = ATT_HEADER_ROW.get(sn, 1)
        headers = read_headers(ws, hrow)
        rows = []
        for r in ws.iter_rows(min_row=hrow + 1, values_only=True):
            # 跳过整行全空的行
            if all(v is None for v in r):
                continue
            rows.append(list(r))
        res[sn] = (headers, rows)
    wb.close()
    return res


def build_mapping(out_headers, att_headers, sheet_name):
    """返回 {out_idx: att_idx} 映射（0-based）。"""
    att_idx = {str(h): i for i, h in enumerate(att_headers)}
    out_to_att = {}
    unmatched_out = []
    for oi, h in enumerate(out_headers):
        key = str(h)
        if key in att_idx:
            out_to_att[oi] = att_idx[key]
        else:
            unmatched_out.append(oi)
    matched_att = set(out_to_att.values())
    unmatched_att = [i for i in range(len(att_headers)) if i not in matched_att]

    if sheet_name in DATE_TO_ORDINAL_SHEETS and unmatched_out and \
       len(unmatched_out) == len(unmatched_att):
        # 按位置映射（月度列）
        for k, oi in enumerate(unmatched_out):
            out_to_att[oi] = unmatched_att[k]
    # 其余未匹配的 out 列保持为空
    return out_to_att


def build_workbook(template_file, out_file, name_map, report):
    # 预读模板所有表头
    wb_t = openpyxl.load_workbook(template_file, read_only=True, data_only=True)
    t_sheets = wb_t.sheetnames
    template_headers = {}
    for sn in t_sheets:
        template_headers[sn] = read_headers(wb_t[sn])
    wb_t.close()

    # 需要的附件 sheet 列表
    att_needed = list(dict.fromkeys(name_map.get(sn, sn) for sn in t_sheets))
    att_data = load_attachment_sheets(att_needed)

    out_wb = openpyxl.Workbook()
    out_wb.remove(out_wb.active)

    for sn in t_sheets:
        a_sn = name_map.get(sn, sn)
        if sn == "说明":
            # 直接复制附件说明表
            if a_sn in att_data:
                ah, arows = att_data[a_sn]
                ws = out_wb.create_sheet(title="说明")
                ws.append(ah)
                blank = 0
                for r in arows:
                    ws.append(r)
                report.append((sn, "说明", len(ah), len(arows), "整表复制(附件输入表说明)"))
            else:
                out_wb.create_sheet(title="说明")
                report.append((sn, "说明", 0, 0, "附件无对应说明表"))
            continue

        if a_sn not in att_data:
            # 附件缺失该表，建立空表头
            ws = out_wb.create_sheet(title=sn)
            ws.append(template_headers[sn])
            report.append((sn, a_sn, len(template_headers[sn]), 0, "附件缺失，仅建表头"))
            continue

        out_headers = template_headers[sn]
        att_headers, att_rows = att_data[a_sn]

        # 初始确认利率曲线：输出列以附件第3行表头为准（含合同组列）
        if sn == "初始确认利率曲线":
            out_headers = att_headers
            out_headers = trim_header(out_headers, sn)

        mapping = build_mapping(out_headers, att_headers, sn)
        ws = out_wb.create_sheet(title=sn)
        ws.append(out_headers)
        n_rows = 0
        n_blank = 0
        for r in att_rows:
            out_row = [None] * len(out_headers)
            has_val = False
            for oi, ai in mapping.items():
                if ai < len(r):
                    v = r[ai]
                    out_row[oi] = v
                    if v is not None:
                        has_val = True
            ws.append(out_row)
            n_rows += 1
            if not has_val:
                n_blank += 1
        note = "列名映射填充"
        if sn in DATE_TO_ORDINAL_SHEETS:
            note = "月度列按位置映射(日期->序数1..60)"
        elif sn == "初始确认利率曲线":
            note = "取附件第3行为表头(含合同组)"
        elif len(mapping) < len(out_headers):
            note = f"模板多列在附件缺失已置空({(len(out_headers)-len(mapping))}列)"
        report.append((sn, a_sn, len(out_headers), n_rows, note))

    out_wb.save(out_file)
    return out_wb


if __name__ == "__main__":
    report = []
    build_workbook(DOCK_T, OUT_DOCK, DOCK_NAME_MAP, report)
    build_workbook(MAN_T, OUT_MAN, MANUAL_NAME_MAP, report)

    print("=" * 100)
    print(f"{'输出sheet':28s} {'附件sheet':28s} {'列数':>5s} {'数据行':>7s}  说明")
    print("=" * 100)
    for sn, a_sn, cols, n, note in report:
        print(f"{sn:28s} {a_sn:28s} {cols:5d} {n:7d}  {note}")
    print("=" * 100)
    print("输出文件:")
    print(" ", OUT_DOCK)
    print(" ", OUT_MAN)
