# -*- coding: utf-8 -*-
import os
from openpyxl import load_workbook
ROOT="D:/IFRS17预测模型"
VBX=os.path.join(ROOT,"验证文件","IFRS17财务预测_改造版_0729_VBA.xlsx")
wb=load_workbook(VBX,data_only=True)
ws=wb["现有业务预期现金流_1"]
print("max_row=",ws.max_row,"max_col=",ws.max_column)
# count rows where 事故期月度序列 (col30, AD) not None
hdr=[c.value for c in ws[1]]
try: adi=hdr.index("事故期月度序列")
except: adi=29
n=0; sample=[]
for ri in range(2,ws.max_row+1):
    v=ws.cell(row=ri,column=adi+1).value
    if v is not None:
        n+=1
        if len(sample)<70: sample.append((ri,int(v)))
print("rows with AD not None:",n)
print("first 70 (row, a):",sample)
# also count 组合 column (col32, AF) non-None
try: afi=hdr.index("组合")
except: afi=31
m=0
for ri in range(2,ws.max_row+1):
    if ws.cell(row=ri,column=afi+1).value is not None: m+=1
print("rows with 组合 not None:",m)
wb.close()
