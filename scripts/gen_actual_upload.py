import json, os, sys, random
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

sys.path.insert(0, os.getcwd())
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'ifrs17_core.settings')
import django; django.setup()
from data_input.views import (
    _PAA_DIRECT_REVENUE_COLS, _PAA_DIRECT_EXPENSE_COLS, _PAA_DIRECT_IFIE_COLS,
    _PAA_CEDING_ALLOC_COLS, _PAA_CEDING_RECOVER_COLS, _PAA_CEDING_IFIE_COLS,
    _to_float, _to_int
)
from collections import defaultdict

random.seed(42)

ACTUAL_DATE = '2026-07-31'
OUTPUT_DIR = r'D:/IFRS17预测模型/验证文件'
OUTPUT_FILE = os.path.join(OUTPUT_DIR, '预实分析_2026-07-31_演示上传.xlsx')

ACTUARIAL_CLASS_NAMES = {
    '32': '交强险', '36': '附加险', '37': '商业三者险', '38': '车损险',
    '101': '企业财产险', '102': '货运险', '104': '责任险', '105': '保证保险',
    '106': '信用保险', '107': '特殊风险', '108': '家财险', '109': '建工险',
    '110': '其他险', '111': '船舶险', '301': '种植业', '302': '养殖业',
    '602': '意外险', '603CC': '商业性健康险', '603ZD': '大病保险',
    '603ED': '大病保险', '603ZX': '其他政策性健康险', '999': '其他',
}

with open('data_input/runtime_cache/calc_results_map.json', 'r', encoding='utf-8') as f:
    mp = json.load(f)
base = mp.get('基础情景') or mp.get('情景0')
fs = base.get('financialStatementsV2') or {}
dates = fs.get('dates') or []
idx = dates.index(ACTUAL_DATE) if ACTUAL_DATE in dates else 1


def perturb(val, rate=0.05, min_abs=None):
    """在 val 附近做小幅随机扰动，保留符号，避免绝对值过小。"""
    if val is None:
        val = 0
    if abs(val) < 1e-6:
        # 对 0 值给一个小的随机非零值（方向随机）
        noise = random.uniform(-max(min_abs or 1000, 100), max(min_abs or 1000, 100))
        return round(noise, 2)
    dev = random.uniform(-rate, rate)
    # 额外小幅固定扰动，确保即使 rate 很小也能看出差异
    fixed = random.uniform(-abs(val) * 0.005, abs(val) * 0.005)
    new_val = val * (1 + dev) + fixed
    return round(new_val, 2)


def perturb_pct(val, rate=3.0):
    """百分比扰动（百分点级别的微调）。"""
    if val is None or abs(val) < 1e-6:
        return round(random.uniform(85, 105), 2)
    delta = random.uniform(-rate, rate)
    return round(val + delta, 2)


# ============================================================
# 1. 公司合计实际数
# ============================================================
ytd = fs.get('ytd', {})

def get_fs_value(section, item_name):
    for r in ytd.get(section, []):
        if r.get('item', '') == item_name:
            vals = r.get('values', [])
            return vals[idx] if idx < len(vals) else 0
    return 0

company_items = [
    ('保险服务收入', get_fs_value('income_statement', '保险服务收入')),
    ('保险服务费用', get_fs_value('income_statement', '保险服务费用')),
    ('投资收益', get_fs_value('income_statement', '投资收益（损失以"-"号填列）')),
    ('营业总收入', get_fs_value('income_statement', '一、营业总收入')),
    ('营业总支出', get_fs_value('income_statement', '二、营业总支出')),
    ('营业利润', get_fs_value('income_statement', '三、营业利润（亏损以"-"号填列）')),
    ('利润总额', get_fs_value('income_statement', '四、利润总额（亏损总额以"-"号填列）')),
    ('净利润', get_fs_value('income_statement', '五、净利润（净亏损以"-"号填列）')),
    ('综合收益总额', get_fs_value('income_statement', '七、综合收益总额')),
    ('承保利润', get_fs_value('income_statement', '承保利润')),
    ('保险合同负债', get_fs_value('balance_sheet', '保险合同负债')),
    ('分出再保险合同资产', get_fs_value('balance_sheet', '分出再保险合同资产')),
    ('资产总计', get_fs_value('balance_sheet', '资产总计')),
    ('负债合计', get_fs_value('balance_sheet', '负债合计')),
]

company_actual = [(item, perturb(val, rate=0.03, min_abs=10000)) for item, val in company_items]

# ============================================================
# 2. 精算险类实际数
# ============================================================
summary = base.get('combinedSummary') or []
raw_groups = defaultdict(list)
for r in summary:
    cls = str(r.get('精算险类') or '未分类')
    if cls in ('2026', 'None', '未分类'):
        continue
    # 合并 .0 后缀
    ccode = cls.replace('.0', '')
    raw_groups[ccode].append(r)

class_rows = []
for ccode in sorted(raw_groups.keys()):
    rows = raw_groups[ccode]
    cname = ACTUARIAL_CLASS_NAMES.get(ccode, ccode)
    # 仅取预测间隔=1（2026-07-31）
    direct_rows = [r for r in rows if str(r.get('业务类型', '')) != '分出' and _to_int(r.get('预测间隔', 1)) == 1]
    ceded_rows = [r for r in rows if str(r.get('业务类型', '')) == '分出' and _to_int(r.get('预测间隔', 1)) == 1]

    ins_rev = ins_exp = ceded_alloc = ceded_recover = ifie = ceded_ifie = 0.0
    for r in direct_rows:
        for col in _PAA_DIRECT_REVENUE_COLS:
            ins_rev += _to_float(r.get(col, 0))
        for col in _PAA_DIRECT_EXPENSE_COLS:
            ins_exp += _to_float(r.get(col, 0))
        for col in _PAA_DIRECT_IFIE_COLS:
            ifie += _to_float(r.get(col, 0))
    for r in ceded_rows:
        for col in _PAA_CEDING_ALLOC_COLS:
            ceded_alloc += _to_float(r.get(col, 0))
        for col in _PAA_CEDING_RECOVER_COLS:
            ceded_recover += _to_float(r.get(col, 0))
        for col in _PAA_CEDING_IFIE_COLS:
            ceded_ifie += _to_float(r.get(col, 0))

    uw_profit = ins_rev - ins_exp - ceded_alloc + ceded_recover - ifie - ceded_ifie
    ceded_net = ceded_recover - ceded_alloc - ceded_ifie
    ifie_net = -ifie - ceded_ifie

    # 小额扰动
    act_ins_rev = perturb(ins_rev, rate=0.04, min_abs=5000)
    act_ins_exp = perturb(ins_exp, rate=0.04, min_abs=5000)
    act_uw_profit = perturb(uw_profit, rate=0.05, min_abs=5000)
    act_ceded_net = perturb(ceded_net, rate=0.05, min_abs=1000)
    act_ifie_net = perturb(ifie_net, rate=0.05, min_abs=1000)

    # 投资收益/其他损益/营业利润/净利润/净资产：按险种规模给一个合理量级的小幅扰动值
    scale = max(abs(act_ins_rev), abs(act_ins_exp), 1)
    invest_income = perturb(scale * 0.02, rate=0.20, min_abs=100)
    other_income = perturb(scale * 0.005, rate=0.30, min_abs=50)
    operating_profit = perturb(act_uw_profit * 0.95, rate=0.04, min_abs=1000)
    net_profit = perturb(operating_profit * 0.75, rate=0.05, min_abs=1000)
    net_asset = perturb(scale * 0.5, rate=0.10, min_abs=1000)

    # 比率：基于扰动后的收支粗略估算
    if abs(act_ins_rev) > 1e-6:
        loss_ratio = round(abs(act_ins_exp) / abs(act_ins_rev) * 100, 2)
        # 费用率给一个围绕 15% 的合理数，但不超过 100
        expense_ratio = round(random.uniform(10, 25), 2)
        combined_ratio = round(loss_ratio + expense_ratio, 2)
    else:
        loss_ratio = expense_ratio = combined_ratio = 0.0

    class_rows.append({
        '评估时点': ACTUAL_DATE,
        '精算险类代码': ccode,
        '精算险类名称': cname,
        '保险服务收入': act_ins_rev,
        '保险服务费用': act_ins_exp,
        '承保利润': act_uw_profit,
        '分出再保险净损益': act_ceded_net,
        '承保财务损益净额': act_ifie_net,
        '投资收益': invest_income,
        '其他损益': other_income,
        '营业利润': operating_profit,
        '净利润': net_profit,
        '净资产': net_asset,
        '综合成本率': combined_ratio,
        '综合赔付率': loss_ratio,
        '综合费用率': expense_ratio,
    })

# ============================================================
# 3. 写入 Excel
# ============================================================
wb = Workbook()

# 公司合计实际数
ws1 = wb.active
ws1.title = '公司合计实际数'
headers1 = ['评估时点', '科目', '实际值']
ws1.append(headers1)
for item, val in company_actual:
    ws1.append([ACTUAL_DATE, item, val])

# 精算险类实际数
ws2 = wb.create_sheet('精算险类实际数')
headers2 = [
    '评估时点', '精算险类代码', '精算险类名称', '保险服务收入', '保险服务费用',
    '承保利润', '分出再保险净损益', '承保财务损益净额', '投资收益', '其他损益',
    '营业利润', '净利润', '净资产', '综合成本率', '综合赔付率', '综合费用率'
]
ws2.append(headers2)
for row in class_rows:
    ws2.append([row[h] for h in headers2])

# 简单样式
header_fill = PatternFill(start_color='E6F4FF', end_color='E6F4FF', fill_type='solid')
header_font = Font(bold=True)
for ws in (ws1, ws2):
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center')
    ws.freeze_panes = 'A2'
    # 自动列宽
    for col in ws.columns:
        max_length = 0
        col_letter = col[0].column_letter
        for cell in col:
            try:
                val_len = len(str(cell.value))
                if val_len > max_length:
                    max_length = val_len
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max(max_length + 2, 10), 30)

os.makedirs(OUTPUT_DIR, exist_ok=True)
wb.save(OUTPUT_FILE)
print('Saved:', OUTPUT_FILE)
print('Company rows:', len(company_actual))
print('Class rows:', len(class_rows))
print('\nCompany actual preview:')
for item, val in company_actual:
    print(f'  {item}: {val:,.2f}')
